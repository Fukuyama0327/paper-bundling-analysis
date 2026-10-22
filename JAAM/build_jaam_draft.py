"""JAAMテンプレートとJAAM/figuresから和文草稿を再生成する。"""

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from io import BytesIO

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'JAAM/2022JAAM_paper_template(J).docx'
OUT = ROOT / 'JAAM/JAAM_橋梁補修契約バンドリング_草稿.docx'
FIG = ROOT / 'JAAM/figures'
doc = Document(SOURCE)

def east_font(run, family, size=None, bold=None):
    run.font.name = family
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    rp = run._element.get_or_add_rPr()
    rf = rp.rFonts
    if rf is None:
        rf = OxmlElement('w:rFonts'); rp.insert(0, rf)
    for a in ('ascii','hAnsi','eastAsia','cs'):
        rf.set(qn('w:'+a), family)
    return run

MINCHO='Hiragino Mincho ProN'
GOTHIC='Hiragino Kaku Gothic ProN'
break_p = next(p for p in doc.paragraphs if p._p.xpath('./w:pPr/w:sectPr'))
body = doc.element.body
for child in list(body):
    if child is break_p._p or child.tag == qn('w:sectPr'):
        continue
    body.remove(child)

# The retained template provides the logo, full-width first section, and page dimensions.
sect0, sect1 = doc.sections
# Use the template's logo only in the first-page header.
logo = sect0.header.part.rels['rId1'].target_part.blob
for s in (sect0,sect1):
    s.different_first_page_header_footer = True
    s.header.is_linked_to_previous = False
    s.first_page_header.is_linked_to_previous = False
    for p in s.header.paragraphs: p.clear()
    for p in s.first_page_header.paragraphs: p.clear()
    s.footer.is_linked_to_previous = False
    s.first_page_footer.is_linked_to_previous = False
    for p in s.footer.paragraphs: p.clear()
    for p in s.first_page_footer.paragraphs: p.clear()
    hp=s.first_page_header.paragraphs[0]
    hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    hp.add_run().add_picture(BytesIO(logo),width=Cm(3.8))
cols = sect1._sectPr.xpath('./w:cols')[0]
cols.set(qn('w:num'),'2'); cols.set(qn('w:space'),'340')

def style(p, size=10, family=MINCHO, bold=False, align=None, before=0, after=0, line=1.12, keep=False):
    p.style = doc.styles['Normal']
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep
    if align is not None: p.alignment = align
    for r in p.runs: east_font(r,family,size,bold)
    return p

def titleline(text, size, after=0, family=MINCHO, bold=False):
    p=break_p.insert_paragraph_before(text)
    style(p,size,family,bold,WD_ALIGN_PARAGRAPH.CENTER,after=after,line=1.12)
    return p

titleline('広域連携に向けた橋梁補修契約の集約効果\n管理エリア設計による定量評価',20,after=12,family=GOTHIC,bold=True)
titleline('福山 峻一¹・水谷 大二郎²',12,after=7)
titleline('¹非会員 東北大学大学院工学研究科（〒980-8579 宮城県仙台市青葉区荒巻字青葉6-6）',9)
titleline('E-mail: fukuyama.shunichi.p3@dc.tohoku.ac.jp',9,after=2)
titleline('²正会員 東北大学大学院工学研究科 准教授（〒980-8579 宮城県仙台市青葉区荒巻字青葉6-6）',9)
titleline('E-mail: daijiro.mizutani.a5@tohoku.ac.jp',9,after=10)
abstract = ('市町村を越えた橋梁の共同発注を検討する際には，どの橋梁を同じ管理エリアで扱えば補修案件を集約しやすいかが問題となる．'
 '本研究は，点検データから推定した年間補修需要確率を用い，地域内橋梁数と一契約の橋梁数上限に応じた年間期待契約件数を算定する．'
 'この指標を最小化する地域分割モデルを宮城県内6市町のRC橋322橋に適用した．管理規模が100～300橋となる範囲で集約効果は大きく高まり，'
 '距離条件に応じて現行の管理区分より期待契約件数が19.4～53.9％低い候補が得られた．本手法は，広域連携の対象範囲を比較するための定量的な材料を示す．')
p=break_p.insert_paragraph_before(abstract); style(p,9,MINCHO,after=7,line=1.12)
p=break_p.insert_paragraph_before('キーワード：橋梁維持管理，広域連携，契約バンドリング，地域分割，期待契約件数')
style(p,9,MINCHO,after=10)

def para(text, after=4):
    p=doc.add_paragraph(text)
    style(p,10,MINCHO,after=after,line=1.16)
    p.paragraph_format.first_line_indent=Cm(0.3)
    return p

def head(text):
    p=doc.add_paragraph(text)
    style(p,11,GOTHIC,True,before=12,after=6,line=1.1,keep=True)
    return p

def sub(text):
    p=doc.add_paragraph(text)
    style(p,10,GOTHIC,True,before=8,after=4,line=1.1,keep=True)
    return p

def formula(text, number):
    p=doc.add_paragraph(text+'     ('+str(number)+')')
    style(p,9,MINCHO,align=WD_ALIGN_PARAGRAPH.CENTER,before=3,after=5,line=1.1)
    return p

def fig(path, caption, width=7.7):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(6)
    p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(FIG/path),width=Cm(width))
    c=doc.add_paragraph(caption)
    style(c,9,MINCHO,align=WD_ALIGN_PARAGRAPH.CENTER,after=7,line=1.1)
    return c

head('1．はじめに')
para('自治体が管理する橋梁の補修需要は，限られた技術職員と予算の下で継続的に生じる．国土交通省が進める地域インフラ群再生戦略マネジメント（群マネ）では，自治体の枠を越える広域連携が選択肢として示されている1)．橋梁補修工事を共同で発注する場合，行政界を越えて補修案件を束ねることにより，契約ごとに繰り返す発注手続を集約できる可能性がある．')
para('その一方，複数市町の管理橋梁を一括して扱えば常に有利とは限らない．同じ年度に補修を要する橋梁がどの程度発生するかは不確実であり，広い地域に橋梁を集めれば現場間の距離も伸びる．どの程度の橋梁数を，どの空間範囲で扱うかを，発注前に比較できる指標が必要である．')
para('本研究では，点検履歴に基づく補修需要の発生確率から，将来の需要を少ない契約単位へ集約できる可能性を「年間期待契約件数」で表す．そして，橋梁間距離に上限を設けた上で，この指標が小さくなる橋梁の管理エリアへの割当を求める．本研究が示すのは共同発注の検討に用いる管理エリア候補と契約集約機会であり，個別工事の契約ロットや施工費そのものを決めるものではない．')

head('2．契約バンドリングと研究の位置付け')
para('米国連邦道路局（FHWA）は，複数の類似事業をまとめて計画・調達する project bundling を，事業実施の時間や費用を抑える手段として紹介している2)．橋梁を対象とする場合には，維持，補修または架替えを要する複数橋を束ねる．Nebraska州の County Bridge Match Program では，単独の郡では束ねる便益を得にくい場合，郡間の協力によって橋梁事業をまとめることが認められている3)．日本の群マネにおける広域連携を考える上でも，行政界より広い範囲に契約の対象を見いだすという発想が参考になる．')
para('既往研究は，劣化に伴う補修需要の予測，所与の事業群の契約バンドリング，地理的単位を組み替える地域分割という三領域で進展してきた4)–6)．本研究はこれらを接続し，補修需要が確率的に発生する段階から，契約集約の機会を評価して管理エリアを設計する．')
para('本研究の貢献は次の三点である．第一に，確率的な補修需要と契約バンドリングを地域分割の評価指標として結び付ける．第二に，点検データから求めた補修需要確率を期待契約件数へ変換し，距離条件の下で橋梁割当を比較する．第三に，実データを用いて，集約効果が高まる管理規模と，行政界を越える割当による改善の大きさを示す．')

head('3．評価方法')
sub('(1) 補修需要と期待契約件数')
para('橋梁の健全度をI，II，IIIの状態で表し，III（IVを含む）への到達を補修需要の発生とみなす．点検間の健全度変化からマルコフ劣化ハザードモデル7)の年間推移確率を推定する．補修後にIへ回復する更新過程を仮定し，単一橋梁に年間補修需要が生じる定常確率をqとする．')
para('ある管理エリアにN橋があり，各橋の補修需要が確率qで独立に生じると仮定すると，年間の補修対象橋梁数Xは二項分布 Binomial(N,q) に従う．同一年度・同一エリア内の補修対象を一契約当たり最大L橋まで束ねられるとき，X=nに必要な契約数は⌈n/L⌉である．従って，年間期待契約件数は次式で求められる．')
formula('f(N,L)＝Σ[n=0…N] Pr(X=n)⌈n/L⌉',1)
para('L=1は個別発注に相当し，L>1では同じ年度に複数需要が生じたときに集約機会が生まれる．Lは実際の契約規模を決定する値ではなく，同時発注を認める範囲を変えたシナリオ条件である．')
sub('(2) 管理エリアの設計')
para('全橋梁をM個の非空管理エリアに一つずつ割り当てる．各エリアkの橋梁数をNₖとし，年間期待契約件数の合計Z＝Σₖf(Nₖ,L)を最小化する．過度な広域化を避けるため，同一エリア内の任意の橋梁間距離が上限D以下となることを制約とする．距離は橋梁座標から求めた大円距離を用いる．')
formula('min Z＝Σₖ f(Nₖ,L)，  エリア内最大距離 ≦ D',2)
para('DとMは実務で選択する条件として複数設定し，得られる期待契約件数と橋梁割当を比較する．距離を費用に換算する重みを置かず，許容距離ごとに結果を示すことで，広域化による集約機会の変化を読み取れるようにした．')

head('4．宮城県6市町への適用')
sub('(1) 対象と補修需要確率')
para('宮城県内のRC橋5,525橋の点検履歴から年間推移確率を推定し，七ヶ宿町，白石市，蔵王町，川崎町，村田町，大河原町の6市町が管理する322橋を地域分割の対象とした8),9)．健全度が改善した観測および30年を超える点検間隔の観測を除外し，IVをIIIに統合した．推定に用いた健全度遷移は7,628件である．')
para('推定した年間推移確率に補修後の状態Iへの回復を組み込み，単一橋梁の年間補修需要確率はq=0.0123（約1.23％）となった．対象6市町ごとに十分な遷移観測を確保できないため，県内RC橋に共通の確率を対象322橋へ適用した．')
fig('study_area.png','図-1　対象6市町のRC橋322橋の分布',7.7)
para('点検間隔の中央値は5年，平均は6.46年であった．制度導入前の点検や，架設時の状態Iから初回点検までの観測を含むため，平均が5年を上回る．こうした観測間隔の違いを考慮して年間推移確率を推定した．補修需要確率qは，いま健全度IIIである橋梁の割合ではなく，補修後に状態Iへ戻る過程を長期的に繰り返したときの年間発生確率である．')
fig('inspection_interval.png','図-2　推定に用いた点検間隔の分布',7.6)
sub('(2) 管理橋梁数と集約機会')
para('まず，管理エリア内橋梁数Nと一契約当たりの橋梁数上限Lによって，集約機会がどう変わるかを調べた．Nが小さい場合，同一年に二橋以上の補修需要が生じる確率は低い．q=0.0123ではN=50橋で12.7％，N=322橋で90.8％である．管理規模の拡大によって同時発注の機会が増えることが分かる．')
fig('expected_contracts_scaling_analysis_a.png','図-3　管理橋梁数と複数補修需要の発生確率',7.6)
para('個別発注（L=1）を基準にした期待契約件数の低減率は，L=5の場合，100～300橋の範囲で大きく高まり，400橋程度以降は改善が緩やかになる．322橋では低減率が69.9％で，理論上限80％に対する到達率は87.4％である．これは対象322橋を一つのエリアで扱った場合の個別発注比であり，現行6市町の管理区分との比較とは異なる．')
fig('expected_contracts_scaling_analysis_b.png','図-4　管理橋梁数と期待契約件数の低減率',7.6)
para('一契約当たりの上限Lが3，5，7，10のいずれでも，理論上限に対する到達率が90％となる管理橋梁数は394～434橋である．L=5では80％到達が226橋，90％到達が397橋であり，効果が現れる規模は数千橋ではなく数百橋のオーダーである．ただし，これらは共通のq，同一年度内の需要集約，工種等によらずL橋まで束ねられるという条件の下での値である．')
fig('expected_contracts_saturation.png','図-5　契約集約効果の理論上限への到達率',7.6)
para('期待契約件数そのものは，管理橋梁数が増えるにつれて増加する．しかしL=1の個別発注では増加がほぼ直線的であるのに対し，L>1では同じエリアで発生する需要をまとめられるため，増加は緩やかになる．図-6はこの差を示す．広域連携の効果を「補修対象橋梁数の減少」と取り違えず，必要な補修をより少ない契約で扱える可能性として読む必要がある．')
fig('expected_contracts_by_bundle_limit.png','図-6　管理橋梁数と一契約の上限別の期待契約件数',7.6)
sub('(3) 距離上限・管理エリア数による比較')
para('主分析ではL=5とし，距離上限Dを15～50 km，管理エリア数Mを1～6で変化させた．現行6市町の管理区分の年間期待契約件数は2.5883件である．最適化した全可解ケースは現行区分より小さく，Dを緩和すると，同じMでの期待契約件数は低下または不変となった．')
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit=False
table.columns[0].width=Cm(3.1); table.columns[1].width=Cm(2.4); table.columns[2].width=Cm(2.1)
for c,t in zip(table.rows[0].cells,['条件（D，M）','期待契約件数','現行比']): c.text=t
rows=[('現行6市町','2.5883','—'),('15 km，6','2.0871','−19.4％'),('20 km，5','1.7912','−30.8％'),('25 km，3','1.5677','−39.4％'),('30 km，3','1.4162','−45.3％'),('35 km，3','1.2620','−51.2％'),('40 km以上，1','1.1933','−53.9％')]
for row in rows:
    cells=table.add_row().cells
    for c,t in zip(cells,row): c.text=t
for row in table.rows:
    for c in row.cells:
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after=Pt(0)
            for r in p.runs: east_font(r,MINCHO,8)
caption=doc.add_paragraph('表-1　距離条件ごとに期待契約件数が最小の解（L=5）')
style(caption,9,MINCHO,align=WD_ALIGN_PARAGRAPH.CENTER,before=5,after=3,line=1.1)
table._tbl.addprevious(caption._p)
para('最も厳しいD=15 kmでも期待契約件数は2.0871件で，現行区分より19.4％低い．D=25 kmでは3エリアで1.5677件，D=35 kmでは3エリアで1.2620件となる．Dが40 km以上なら322橋を一つのエリアに含められ，1.1933件，現行比53.9％の低減となる．この最大値は40 km以上・1エリアという条件付きの結果である．')
fig('dm_sensitivity.png','図-7(a)　距離上限と管理エリア数別の期待契約件数',7.7)
fig('dm_sensitivity_b.png','図-7(b)　距離上限40 km以上の拡大',7.7)
sub('(4) 橋梁割当からみた検討材料')
para('図-8には，D=25 km，M=3の橋梁割当を示す．形成された三つのエリアは277橋，35橋，10橋からなり，現行の市町境界とは一致しない．行政界を越えて近接する橋梁を同じ候補エリアに含めることで，単独の市町では限られる同時発注機会を広げられる．')
fig('districting_map_D25_M3.png','図-8　D=25 km，M=3における橋梁の割当',7.7)
para('一方，結果は管理エリアごとの橋梁数が大きく偏ることも示す．D=25 km，M=3の10橋エリアは，277橋エリアより集約機会が限られる．D=30 km，M=2の例では251橋と71橋の二つのエリアが形成され，期待契約件数はそれぞれ1.0484件，0.5858件となる．これらの例は，地域数だけでなく，どの橋梁を各エリアに含めるかが重要であることを示す．')
fig('region_breakdown.png','図-9　代表的な管理エリア案の橋梁数と期待契約件数',7.6)
para('D=25 km，M=3の期待契約件数は1.5677件である．これに対し，D=30 km，M=2の上記の割当では1.6342件となり，エリア数が少なく距離上限が緩くても，常に期待契約件数が小さくなるわけではない．二つの条件が同時に異なるためであり，実務上もエリア数と許容距離を組み合わせて候補案を比較することが適切である．図中の色分け面は橋梁割当を見やすくするためのVoronoi表現であり，実際の管理区域境界ではない．')
para('図-10はD=30 km，M=2の橋梁割当を示す．西部から南部にかけての71橋と，残る251橋が二つのエリアとなる．D=25 km，M=3の解と比較すると，行政界を越える組合せ方だけでなく，少数の橋梁が残るエリアの位置も変わる．このように候補案を地図と期待契約件数の両方で提示すれば，参加自治体が実際の施工範囲や管理体制を議論しやすくなる．')
fig('districting_map_D30_M2.png','図-10　D=30 km，M=2における橋梁の割当',7.7)
para('距離上限Dを緩和したときの期待契約件数の改善は，現場の移動や契約調整の負担を差し引いた純便益ではない．ただし，Dを厳しくした候補から順に比較すれば，追加の広域化によってどの程度の集約機会が得られるかを確認できる．その上で，各候補の工種，発注時期，地元企業の参加可能性等を調べるという段階的な検討が考えられる．')
sub('(5) 広域連携案を比較する際の視点')
para('管理エリア数Mは，期待契約件数だけから定めるべき値ではない．Mを少なくすれば個々のエリアにより多くの橋梁を含められるが，共同発注の契約主体や費用分担を調整する範囲も変わる．したがって，連携可能な市町の組合せを想定してMを設定し，その条件ごとに距離上限Dを変えた結果を比較するのが自然である．図-7の曲線は，この条件付きの比較を示している．')
para('現行の管理区分と比べる際には，まず現行6市町の橋梁数を同じ補修需要確率qと同じ上限Lで評価する．そのうえで最適化案の期待契約件数との差をみる．表-1の19.4～53.9％は，この共通条件の下で算定した差であり，実際の発注件数を予測した値ではない．現行区分は距離上限Dを課して選ばれた案でもないため，表の各行を「同じ地理条件下の改善率」とは解釈しない．')
para('次に，候補案の地図を確認する．本モデルは橋梁を点として扱い，同一エリア内の橋梁間の最大大円距離を制約する．この値は車両の走行距離や移動時間を直接表さない．図-8と図-10の割当を検討する際には，道路網での移動，拠点からの到達，山地や河川による分断を確認し，現場の実態と合わない橋梁の組合せを見直す必要がある．')
para('最後に，同一年度に補修需要が生じても，工種，施工条件，予算や発注時期が揃わなければ一契約にまとめられない．本研究のLはこうした条件を詳細に再現したものではなく，何橋まで同時発注を許容するかを変えた感度分析である．したがって本手法は，契約仕様を直接作成するものではなく，どの管理範囲を優先して共同発注の詳細検討へ進めるかを選ぶために使う．')

head('5．おわりに')
para('本研究は，点検データに基づく確率的な補修需要を年間期待契約件数に変換し，距離条件の下で行政界に依らない橋梁の管理エリア候補を比較した．宮城県の事例では，集約効果は管理規模100～300橋の範囲で大きく高まり，6市町の現行区分より期待契約件数が小さい候補が得られた．距離上限を厳しく設定した場合にも改善余地があることは，広域連携の範囲を段階的に検討する材料となる．')
para('期待契約件数は，契約単位で繰り返す発注手続の集約機会を表す指標である．実際の共同発注では，工種や施工時期の整合，現場間の道路移動，地元企業の参加可能性，費用分担と契約主体の調整を併せて検討する必要がある．本研究の橋梁割当は，それらの条件を検討する出発点として用いることができる．')

head('参考文献')
refs=[
'1) 国土交通省：地域インフラ群再生戦略マネジメント手引き Ver.1，2025．',
'2) FHWA: Project Bundling Fact Sheet (EDC-5), Federal Highway Administration, 2020.',
'3) FHWA: Bridge Bundling Guidebook, Federal Highway Administration, 2019.',
'4) Qiao, J. Y. et al.: A proposed algorithm for identifying heuristic-optimal bundling strategies, Computer-Aided Civil and Infrastructure Engineering, 2026.',
'5) Kalcsics, J. and Rios-Mercado, R. Z.: Districting Problems, Location Science, pp.705–743, 2019.',
'6) Nakazato, Y., Mizutani, D. and Fukuyama, S.: Optimal Repair Policies for Infrastructure Systems, Journal of Infrastructure Systems, 29(3), 2023.',
'7) 津田尚胤ほか：橋梁劣化予測のためのマルコフ推移確率の推定，土木学会論文集，No.801/I-73，pp.69–82，2005．',
'8) 全国道路施設点検データベース，https://road-structures-db.mlit.go.jp．',
'9) 国土交通省：道路メンテナンス年報，平成26年度～令和6年度．',
]
for t in refs:
    p=doc.add_paragraph(t)
    style(p,8,MINCHO,after=2,line=1.05)
    p.paragraph_format.left_indent=Cm(0.35); p.paragraph_format.first_line_indent=Cm(-0.35)

# A final continuous section makes Word balance the two columns on the last page.
tail=doc.add_section(WD_SECTION_START.CONTINUOUS)
tail.header.is_linked_to_previous=False
tail.first_page_header.is_linked_to_previous=False
tail.footer.is_linked_to_previous=False
tail.first_page_footer.is_linked_to_previous=False
for p in tail.header.paragraphs+tail.first_page_header.paragraphs+tail.footer.paragraphs+tail.first_page_footer.paragraphs:
    p.clear()

doc.save(OUT)
print(OUT)
print('abstract characters',len(abstract))
